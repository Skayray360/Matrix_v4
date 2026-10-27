# Creado por Aldo Garcia.
"""Conservar estructura y procedencia usando documentos sinteticos reales."""

import io

import docx
import pytest
from openpyxl import Workbook

from app.ingestion.loaders import extract_document
from app.rag.chunking import chunk_blocks, estimate_tokens
from tests.unit.test_pdf_rag_extraction import mixed_pdf

pytestmark = pytest.mark.unit


def workbook_bytes(workbook):
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def test_excel_chunks_keep_rows_labels_and_physical_row_numbers():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Personal"
    sheet.append(["Puesto", "Descripcion", "Dias"])
    sheet.append([None, None, None])
    for number in range(40):
        sheet.append([f"Puesto{number}", "Primera linea\nsegunda linea", number + 12])
    extracted = extract_document(workbook_bytes(workbook), filename="personal.xlsx")
    chunks = chunk_blocks(extracted.blocks, chunk_size_tokens=80, overlap_tokens=10)
    assert len(chunks) > 5
    assert all(estimate_tokens(chunk.text) <= 80 and chunk.page_or_sheet == "Personal" for chunk in chunks)
    for number in range(40):
        expected = (f"- Fila {number + 3}: Puesto: Puesto{number}; "
                    f"Descripcion: Primera linea segunda linea; Dias: {number + 12}")
        assert sum(expected in chunk.text.splitlines() for chunk in chunks) == 1


def test_excel_formula_without_cached_value_is_reported_without_execution():
    workbook = Workbook()
    workbook.active.append(["Resultado"])
    workbook.active.append(["=SUM(10,20)"])
    extracted = extract_document(workbook_bytes(workbook), filename="formula.xlsx")
    assert any("1 formulas sin resultado guardado" in warning for warning in extracted.warnings)
    assert "SUM" not in extracted.plain_text()


def test_empty_workbook_does_not_become_evidence_from_a_sheet_title():
    extracted = extract_document(workbook_bytes(Workbook()), filename="empty.xlsx")
    assert extracted.is_empty and extracted.warnings


def test_oversized_excel_row_fails_instead_of_detaching_values_from_labels():
    workbook = Workbook()
    workbook.active.append(["Condiciones"])
    workbook.active.append(["condicion " * 100])
    extracted = extract_document(workbook_bytes(workbook), filename="wide.xlsx")
    with pytest.raises(ValueError, match="fila"):
        chunk_blocks(extracted.blocks, chunk_size_tokens=60, overlap_tokens=10)


@pytest.mark.parametrize("outer_pipes", [False, True])
def test_markdown_tables_repeat_headers_and_keep_complete_rows(outer_pipes):
    lines = ["Puesto | Dias", "--- | ---", *[f"Puesto{i} | {i + 12}" for i in range(30)]]
    if outer_pipes:
        lines = [f"| {line} |" for line in lines]
    source = "# Vacaciones\n\n" + "\n".join(lines)
    extracted = extract_document(source.encode(), filename="politica.md")
    assert len([block for block in extracted.blocks if block.kind == "table"]) == 1
    chunks = chunk_blocks(extracted.blocks, chunk_size_tokens=65, overlap_tokens=10)
    table_chunks = [chunk for chunk in chunks if "Puesto0" in chunk.text or "| Puesto" in chunk.text]
    assert len(table_chunks) > 1
    assert all("| Puesto | Dias |\n| --- | --- |" in chunk.text for chunk in table_chunks)
    for number in range(30):
        assert sum(f"| Puesto{number} | {number + 12} |" in chunk.text.splitlines() for chunk in chunks) == 1


@pytest.mark.parametrize("fence", ["```", "~~~~"])
def test_markdown_code_keeps_sections_indentation_and_fences(fence):
    source = f"# Manual\n{fence}python\n# Comentario, no seccion\n"
    source += "\n".join(f"    valor_{number} = {number}" for number in range(30))
    source += f"\n{fence}\n\nTexto final."
    extracted = extract_document(source.encode(), filename="manual.md")
    code = next(block for block in extracted.blocks if block.kind == "code")
    assert code.section == "Manual" and "    valor_0 = 0" in code.text
    assert all(block.section == "Manual" for block in extracted.blocks)
    chunks = chunk_blocks([code], chunk_size_tokens=40, overlap_tokens=5)
    assert len(chunks) > 1
    assert all(chunk.text.startswith(fence + "python\n") and chunk.text.endswith(fence) for chunk in chunks)
    assert all(estimate_tokens(chunk.text) <= 40 for chunk in chunks)


def test_markdown_escaped_pipes_and_code_table_examples_are_literal():
    source = "A | B\n--- | ---\none \\| two | three\n\n```md\nA | B\n--- | ---\n```"
    extracted = extract_document(source.encode(), filename="escaped.md")
    table = next(block for block in extracted.blocks if block.kind == "table")
    assert "| one \\| two | three |" in table.text
    code = next(block for block in extracted.blocks if block.kind == "code")
    assert code.text == "```md\nA | B\n--- | ---\n```"


def test_word_table_becomes_markdown_with_original_table_locator():
    document = docx.Document()
    document.add_heading("Vacaciones", level=1)
    table = document.add_table(rows=2, cols=2)
    for row, cells in zip(table.rows, [("Antiguedad", "Dias"), ("1", "12")], strict=True):
        for cell, text in zip(row.cells, cells, strict=True):
            cell.text = text
    buffer = io.BytesIO()
    document.save(buffer)
    extracted = extract_document(buffer.getvalue(), filename="politica.docx")
    block = next(block for block in extracted.blocks if block.kind == "table")
    assert block.text == "| Antiguedad | Dias |\n| --- | --- |\n| 1 | 12 |"
    assert block.page_or_sheet == "tabla 1" and block.section == "Vacaciones"


@pytest.mark.parametrize("extension", [".docx", ".xlsx", ".pdf", ".md"])
def test_supported_formats_roundtrip_through_the_real_isolated_worker(extension):
    from app.ingestion.isolated_extraction import extract_document as isolated_extract

    if extension == ".pdf":
        data = mixed_pdf(image=False)
        expected = "Prima del 7.5%"
    elif extension == ".xlsx":
        workbook = Workbook()
        workbook.active.append(["Puesto", "Dias"])
        workbook.active.append(["Tecnico", 12])
        data = workbook_bytes(workbook)
        expected = "- Fila 2: Puesto: Tecnico; Dias: 12"
    elif extension == ".docx":
        document = docx.Document()
        document.add_heading("Vacaciones", level=1)
        document.add_paragraph("Tecnico: 12 dias")
        buffer = io.BytesIO()
        document.save(buffer)
        data = buffer.getvalue()
        expected = "Tecnico: 12 dias"
    else:
        data = b"# Manual\n\n```python\n# Comentario\nvalor = 12\n```"
        expected = "```python\n# Comentario\nvalor = 12\n```"
    extracted = isolated_extract(data, filename="synthetic" + extension)
    chunks = chunk_blocks(extracted.blocks, chunk_size_tokens=900, overlap_tokens=120)
    assert chunks and any(expected in chunk.text for chunk in chunks)
    if extension in {".pdf", ".docx", ".xlsx"}:
        assert all(chunk.page_or_sheet for chunk in chunks)
