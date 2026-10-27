# Creado por Aldo Garcia.
"""Cobertura de la fuente, distinta de recuperar todos los chunks disponibles."""

from io import BytesIO
from types import SimpleNamespace

import docx
import pytest
from docx.oxml import OxmlElement
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference

from app.database.models import Document
from app.ingestion import isolated_extraction, loaders
from app.ingestion.local_ocr import OcrOptions
from app.rag.retriever import Retriever
from tests.conftest import make_context
from tests.unit import test_document_overview_retrieval_20261008 as summary_fixtures
from tests.unit import test_private_reindex_final as private_fixtures
from tests.unit.test_multiformat_ingestion import presentation
from tests.unit.test_pdf_rag_extraction import mixed_pdf

pytestmark = pytest.mark.unit
environment = private_fixtures.environment
corpus = summary_fixtures.corpus


def test_csv_truncation_is_structured_and_preserves_available_rows(monkeypatch):
    data = b"Name,Amount\nA,10\nB,20\n"
    assert not loaders.extract_csv(data).extraction_incomplete
    monkeypatch.setattr(loaders, "MAX_CSV_ROWS", 2)
    result = loaders.extract_csv(data)
    assert result.extraction_incomplete and "Amount: 10" in result.plain_text()
    assert "Amount: 20" not in result.plain_text()


def test_csv_values_without_a_header_do_not_silently_accredit_complete_coverage():
    assert not loaders.extract_csv(b"Name,Amount\nA,10,\n").extraction_incomplete
    result = loaders.extract_csv(b"Name,Amount\nA,10,Additional rule\n")
    assert "Amount: 10" in result.plain_text()
    assert "Additional rule" not in result.plain_text()
    assert result.extraction_incomplete and result.warnings


@pytest.mark.parametrize("limitation", ["sheets", "rows", "columns", "formula"])
def test_xlsx_omissions_have_structured_coverage(limitation, monkeypatch):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Name", "Amount", "Currency"])
    sheet.append(["A", 10, "MXN"])
    sheet.append(["B", 20, "MXN"])
    if limitation == "sheets":
        workbook.create_sheet("Unprocessed").append(["Extra", "Data"])
        monkeypatch.setattr(loaders, "MAX_XLSX_SHEETS", 1)
    elif limitation == "rows":
        monkeypatch.setattr(loaders, "MAX_XLSX_ROWS_PER_SHEET", 2)
    elif limitation == "columns":
        monkeypatch.setattr(loaders, "MAX_XLSX_COLUMNS", 2)
    else:
        sheet["B3"] = "=10+10"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    result = loaders.extract_xlsx(output.getvalue())
    assert result.extraction_incomplete and not result.is_empty


def test_xlsx_chart_text_is_not_counted_as_extracted_cells():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Policy", "Amount"])
    sheet.append(["Synthetic", 10])
    output = BytesIO()
    workbook.save(output)
    assert not loaders.extract_xlsx(output.getvalue()).extraction_incomplete
    chart = BarChart()
    chart.title = "Additional chart rule."
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=2), titles_from_data=True)
    sheet.add_chart(chart, "D1")
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    result = loaders.extract_xlsx(output.getvalue())
    assert "Amount: 10" in result.plain_text()
    assert "Additional chart rule." not in result.plain_text()
    assert result.extraction_incomplete and result.warnings


def test_pdf_ocr_warning_is_not_itself_proof_of_missing_content(monkeypatch):
    assert not loaders.extract_pdf(mixed_pdf(image=False)).extraction_incomplete
    assert loaders.extract_pdf(mixed_pdf()).extraction_incomplete
    monkeypatch.setattr(loaders, "ocr_pdf_page", lambda *_args, **_kwargs: "Additional image text.")
    complete = loaders.extract_pdf(mixed_pdf(), ocr_options=OcrOptions(enabled=True))
    assert complete.warnings and not complete.extraction_incomplete
    monkeypatch.setattr(loaders, "ocr_pdf_page", lambda *_args, **_kwargs: "")
    assert loaders.extract_pdf(mixed_pdf(), ocr_options=OcrOptions(enabled=True)).extraction_incomplete


def test_isolated_parser_transports_coverage_without_exposing_document_text():
    result = isolated_extraction.extract_document(mixed_pdf(), filename="synthetic.pdf")
    assert result.extraction_incomplete and not result.is_empty


def test_pptx_missing_text_and_complete_slides_are_distinguished():
    assert loaders.extract_pptx(presentation((1,), empty=True)).extraction_incomplete
    assert not loaders.extract_pptx(presentation((1,))).extraction_incomplete


def test_docx_embedded_drawing_is_not_counted_as_extracted_text():
    document = docx.Document()
    paragraph = document.add_paragraph("Available synthetic policy.")
    output = BytesIO()
    document.save(output)
    assert not loaders.extract_docx(output.getvalue()).extraction_incomplete
    paragraph.add_run()._r.append(OxmlElement("w:drawing"))
    output = BytesIO()
    document.save(output)
    result = loaders.extract_docx(output.getvalue())
    assert result.extraction_incomplete and "Available synthetic policy." in result.plain_text()


@pytest.mark.parametrize("story", ["header", "footer"])
def test_docx_header_and_footer_text_is_not_counted_as_extracted_body(story):
    document = docx.Document()
    document.add_paragraph("Available synthetic policy.")
    getattr(document.sections[0], story).paragraphs[0].text = "Additional story rule."
    output = BytesIO()
    document.save(output)
    result = loaders.extract_docx(output.getvalue())
    assert "Available synthetic policy." in result.plain_text()
    assert "Additional story rule." not in result.plain_text()
    assert result.extraction_incomplete and result.warnings


@pytest.mark.parametrize("scope", ["corporate", "conversation"])
@pytest.mark.parametrize("coverage", ["complete", "partial", "legacy"])
def test_summary_does_not_hide_extraction_gaps_even_when_all_chunks_fit(corpus, scope, coverage):
    chunks = corpus.publish("Synthetic.docx", scope_name=scope)
    client = corpus.store._client
    arguments = {"collection_name": corpus.store.collection_for(scope),
                 "points": [chunk.metadata.chunk_id for chunk in chunks]}
    if coverage == "legacy":
        client.delete_payload(keys=["extraction_incomplete"], **arguments)
    else:
        client.set_payload(payload={"extraction_incomplete": coverage == "partial"}, **arguments)
    if scope == "corporate":
        result = corpus.retriever.retrieve(
            ctx=corpus.ctx, question="Resume Synthetic.docx", authorized_categories=frozenset({"tecnologia"}),
            summary=True, include_private=False,
        )
    else:
        result = corpus.retriever.retrieve_attachment_summary(
            ctx=corpus.ctx, conversation_id="thread", question="Resume Synthetic.docx",
        )
    assert len(result.evidences) == len(chunks)
    assert result.truncated is (coverage != "complete")
    query = Retriever(store=corpus.store, llm=SimpleNamespace(embed_one=lambda _text: [1.0, 0.0, 0.0]))
    factual = query.retrieve(
        ctx=corpus.ctx, question="Synthetic.docx", authorized_categories=frozenset({"tecnologia"}),
        conversation_id="thread" if scope == "conversation" else None,
    )
    assert factual.evidences and factual.truncated is (coverage != "complete")
    if scope == "conversation":
        legacy_path = corpus.retriever.retrieve_attachment_summary(ctx=corpus.ctx, conversation_id="thread")
        assert legacy_path.truncated is (coverage != "complete")
        assert not corpus.retriever.retrieve_attachment_summary(
            ctx=make_context(user_id="another-owner"), conversation_id="thread",
        ).evidences


def test_partial_attachment_is_published_once_and_historical_metadata_can_be_upgraded(environment, monkeypatch):
    env = environment
    monkeypatch.setattr(loaders, "MAX_CSV_ROWS", 2)
    data = b"Name,Amount\nA,10\nB,20\n"
    prepared = env.service.prepare_attachment(data=data, filename="synthetic.csv")
    assert prepared.extraction_incomplete
    outcome = env.service.ingest_conversation_attachment(
        env.db, data=data, display_name="synthetic.csv", internal_filename="synthetic.csv",
        mime_type="text/csv", owner_user_id="user-a", conversation_id="chat-a", prepared=prepared,
    )
    env.db.commit()
    document = env.db.get(Document, outcome.document_id)
    assert env.service._document_index_is_complete(document)
    assert env.service.reindex_conversation_attachment(env.db, document=document).skipped
    result = private_fixtures.summary(env)
    assert result.evidences and result.truncated
    collection = env.store.collection_for("conversation")
    points, _ = env.store._client.scroll(collection_name=collection, limit=10)
    env.store._client.delete_payload(collection_name=collection, keys=["extraction_incomplete"],
                                    points=[point.id for point in points])
    assert not env.service._document_index_is_complete(document)
    updated = env.service.reindex_conversation_attachment(env.db, document=document)
    env.db.commit()
    assert not updated.skipped and env.service._document_index_is_complete(document)
    assert private_fixtures.summary(env).truncated
